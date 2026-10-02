from __future__ import annotations

import unittest
import openpyxl
from io import BytesIO
from core.excel_manager import ExcelPricingManager


class TestExcelPricingManager(unittest.TestCase):
    def setUp(self):
        self.sample_matrix = {
            "iPhone 16 Pro Max": {
                "256": {"max_buy": 105000, "market": 125000, "enabled": True},
                "512": {"max_buy": 118000, "market": 140000, "enabled": False},
            },
            "iPhone 15": {
                "128": {"max_buy": 45000, "market": 55000, "enabled": True},
            },
        }

    def test_export_and_read_back(self):
        # Экспорт в байты
        excel_bytes = ExcelPricingManager.export_matrix_to_bytes(self.sample_matrix)
        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 1000)

        # Открываем книгу openpyxl для проверки структуры
        wb = openpyxl.load_workbook(BytesIO(excel_bytes))
        ws = wb.active
        self.assertEqual(ws.title, "Лимиты Выкупа Apple")

        # Проверяем заголовки
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers, ["Модель iPhone", "Память (GB)", "Лимит выкупа (₽)", "Рыночная цена (₽)", "Статус"])

        # Проверяем строки данных
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 3)

        # Проверяем значения
        self.assertEqual(rows[0], ("iPhone 16 Pro Max", 256, 105000, 125000, "ВКЛЮЧЕН"))
        self.assertEqual(rows[1], ("iPhone 16 Pro Max", 512, 118000, 140000, "ВЫКЛЮЧЕН"))
        self.assertEqual(rows[2], ("iPhone 15", 128, 45000, 55000, "ВКЛЮЧЕН"))

    def test_import_modified_excel(self):
        # Создаем Excel с изменениями: меняем цену и статус
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Лимиты Выкупа Apple"
        ws.append(["Модель iPhone", "Память (GB)", "Лимит выкупа (₽)", "Рыночная цена (₽)", "Статус"])
        ws.append(["iPhone 16 Pro Max", "256", 110000, 130000, "ВКЛ"])
        ws.append(["iPhone 16 Pro Max", "512", 125000, 145000, "ВКЛ"])  # был ВЫКЛ, стал ВКЛ
        ws.append(["iPhone 15", "128", 42000, 52000, "ВЫКЛ"])  # был ВКЛ, стал ВЫКЛ

        buf = BytesIO()
        wb.save(buf)
        excel_bytes = buf.getvalue()

        # Импортируем без сохранения на диск (dry run)
        updated_matrix, count, errors = ExcelPricingManager.parse_excel_to_matrix(excel_bytes)
        self.assertEqual(count, 3)
        self.assertEqual(len(errors), 0)

        # Проверяем обновленные данные
        self.assertEqual(updated_matrix["iPhone 16 Pro Max"]["256"]["max_buy"], 110000)
        self.assertEqual(updated_matrix["iPhone 16 Pro Max"]["256"]["market"], 130000)
        self.assertTrue(updated_matrix["iPhone 16 Pro Max"]["256"]["enabled"])

        self.assertEqual(updated_matrix["iPhone 16 Pro Max"]["512"]["max_buy"], 125000)
        self.assertTrue(updated_matrix["iPhone 16 Pro Max"]["512"]["enabled"])

        self.assertEqual(updated_matrix["iPhone 15"]["128"]["max_buy"], 42000)
        self.assertFalse(updated_matrix["iPhone 15"]["128"]["enabled"])

    def test_import_invalid_data_handles_cleanly(self):
        # Пустой или поврежденный файл не должен ронять приложение
        count, errors = ExcelPricingManager.import_matrix_from_bytes(b"not an excel file")
        self.assertEqual(count, 0)
        self.assertGreater(len(errors), 0)


if __name__ == "__main__":
    unittest.main()
