from __future__ import annotations

import io
import json
import logging
from pathlib import Path
from typing import Dict, Any, Tuple, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class ExcelPricingManager:
    """
    Менеджер экспорта и импорта матрицы пороговых цен через Excel (.xlsx).
    Позволяет перекупщику настраивать лимиты выкупа прямо с телефона/ПК в таблице.
    """

    @classmethod
    def export_matrix_to_bytes(cls, matrix: dict[str, dict[str, dict[str, Any]]]) -> bytes:
        """Генерирует красиво оформленный Excel-файл из матрицы цен."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Лимиты Выкупа Apple"

        # Стили
        header_fill = PatternFill(start_color="1E3A1E", end_color="1E3A1E", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Arial", size=10)
        bold_font = Font(name="Arial", size=10, bold=True)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0"),
        )

        headers = [
            "Модель iPhone",
            "Память (GB)",
            "Лимит выкупа (₽)",
            "Рыночная цена (₽)",
            "Статус",
        ]

        ws.append(headers)
        ws.row_dimensions[1].height = 28

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Заполнение строк
        row_num = 2
        for model, storages in matrix.items():
            for storage, data in storages.items():
                max_buy = data.get("max_buy", 0)
                market = data.get("market", 0)
                enabled = data.get("enabled", True)
                status_text = "ВКЛЮЧЕН" if enabled else "ВЫКЛЮЧЕН"

                ws.append([
                    model,
                    int(storage) if str(storage).isdigit() else str(storage),
                    max_buy,
                    market,
                    status_text,
                ])

                ws.row_dimensions[row_num].height = 20

                # Оформление ячеек строки
                ws.cell(row=row_num, column=1).alignment = left_align
                ws.cell(row=row_num, column=1).font = bold_font

                ws.cell(row=row_num, column=2).alignment = center_align
                ws.cell(row=row_num, column=2).font = regular_font

                for c_idx in (3, 4):
                    c = ws.cell(row=row_num, column=c_idx)
                    c.alignment = right_align
                    c.font = regular_font
                    c.number_format = '#,##0 "₽"'

                status_cell = ws.cell(row=row_num, column=5)
                status_cell.alignment = center_align
                status_cell.font = bold_font
                if enabled:
                    status_cell.font = Font(name="Arial", size=10, bold=True, color="008000")
                else:
                    status_cell.font = Font(name="Arial", size=10, bold=True, color="808080")

                for c_idx in range(1, 6):
                    ws.cell(row=row_num, column=c_idx).border = thin_border

                row_num += 1

        # Автоподбор ширины колонок
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    @classmethod
    def parse_excel_to_matrix(
        cls, file_bytes: bytes
    ) -> tuple[dict[str, dict[str, dict[str, Any]]], int, list[str]]:
        """
        Парсит байты Excel таблицы и преобразует их в структуру матрицы цен.
        Возвращает:
            (матрица, количество_загруженных_конфигураций, список_ошибок_или_предупреждений)
        """
        errors: list[str] = []
        new_matrix: dict[str, dict[str, dict[str, Any]]] = {}
        count = 0

        try:
            buf = io.BytesIO(file_bytes)
            wb = openpyxl.load_workbook(buf, data_only=True)
            ws = wb.active
        except Exception as e:
            return {}, 0, [f"Не удалось открыть Excel файл: {e}"]

        if ws is None:
            return {}, 0, ["В таблице отсутствует активный лист."]

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not any(row):
                continue

            if not row[0]:
                continue

            model = str(row[0]).strip()
            if not model:
                continue

            raw_storage = str(row[1]).strip() if len(row) > 1 and row[1] is not None else "128"
            storage = "".join(filter(str.isdigit, raw_storage)) or "128"

            try:
                max_buy = int(float(str(row[2]).replace(" ", "").replace("₽", "").replace(",", "."))) if len(row) > 2 and row[2] is not None else 0
            except (ValueError, TypeError):
                errors.append(f"Строка {row_idx}: некорректная цена выкупа '{row[2]}'")
                max_buy = 0

            try:
                market = int(float(str(row[3]).replace(" ", "").replace("₽", "").replace(",", "."))) if len(row) > 3 and row[3] is not None else max_buy
            except (ValueError, TypeError):
                market = max_buy

            enabled = True
            if len(row) > 4 and row[4] is not None:
                status_str = str(row[4]).strip().upper()
                if any(x in status_str for x in ["ВЫКЛ", "НЕТ", "OFF", "0", "FALSE", "ВЫКЛЮЧЕН"]):
                    enabled = False

            if model not in new_matrix:
                new_matrix[model] = {}

            new_matrix[model][storage] = {
                "max_buy": max_buy,
                "market": market,
                "enabled": enabled,
            }
            count += 1

        if not new_matrix and not errors:
            errors.append("В таблице не найдено ни одной строки с ценами.")

        return new_matrix, count, errors

    @classmethod
    def import_matrix_from_bytes(
        cls, file_bytes: bytes, target_path: Optional[str] = None
    ) -> tuple[int, list[str]]:
        """
        Парсит загруженный пользователем .xlsx файл, валидирует и сохраняет в JSON.
        Возвращает:
            (количество_загруженных_конфигураций, список_ошибок)
        """
        new_matrix, count, errors = cls.parse_excel_to_matrix(file_bytes)
        if count == 0:
            return 0, errors

        # Сохранение в pricing_matrix.json
        from config import settings
        pricing_file = Path(target_path) if target_path else (settings.BASE_DIR / settings.PRICING_FILE)
        try:
            with open(pricing_file, "w", encoding="utf-8") as f:
                json.dump(new_matrix, f, ensure_ascii=False, indent=2)
            logger.info("Матрица цен обновлена из Excel: %d конфигураций сохранено в %s", count, pricing_file)
        except Exception as e:
            err_msg = f"Ошибка при сохранении JSON матрицы: {e}"
            logger.error(err_msg)
            errors.append(err_msg)
            return 0, errors

        return count, errors
